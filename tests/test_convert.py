"""Excel and PowerPoint exports: the detail that makes them usable."""
import io, os, re, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pymupdf as fitz
from openpyxl import load_workbook
from pptx import Presentation

from backend import convert

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
SRC = os.path.join(OUT, "fixture.pdf")
HARD = os.path.join(OUT, "hard.pdf")
findings = []


def check(label, got, want):
    ok = want(got) if callable(want) else got == want
    if not ok:
        findings.append((label, got))
    print("  %s %-46s %r" % ("ok  " if ok else "BUG ", label, got))


print("== excel ==")
for raw, expect in [("$100,000", 100000.0), ("(1,234)", -1234.0), ("12.5%", 0.125),
                    ("1,234.56", 1234.56), ("-7", -7.0)]:
    check("reads %r as a number" % raw, convert._coerce(raw)[0], expect)
for raw in ["Jan 15", "Deposit", "N/A", "4.2.1", ""]:
    check("leaves %r as text" % raw, convert._coerce(raw)[0], raw)

result = convert.to_xlsx(fitz.open(SRC), os.path.join(OUT, "c_fixture.xlsx"))
book = load_workbook(os.path.join(OUT, "c_fixture.xlsx"))
sheet = book[book.sheetnames[0]]
numbers = [c.value for row in sheet.iter_rows(min_row=2) for c in row
           if isinstance(c.value, (int, float))]
check("amounts land as numbers", numbers, [100000, 4150000])
check("so the column can be summed", sum(numbers), 4250000)
check("currency formatting is kept",
      sheet.cell(row=2, column=2).number_format, lambda f: "$" in f)
check("header row is frozen", sheet.freeze_panes, "A2")
check("filter is set", bool(sheet.auto_filter.ref), True)
check("sheet is named after the table or page", sheet.title,
      lambda t: t and len(t) <= 31)


class FakeBook:
    sheetnames = ["Page 1 table 1"]


check("illegal sheet characters are stripped",
      convert._sheet_name(FakeBook(), "Rates [2026]: fees/levies"),
      lambda n: not set(n) & set("[]:*?/" + chr(92)))
check("duplicate sheet names are made unique",
      convert._sheet_name(FakeBook(), "Page 1 table 1"),
      lambda n: n != "Page 1 table 1")

no_tables = convert.to_xlsx(fitz.open(HARD), os.path.join(OUT, "c_hard.xlsx"))
check("a document with no tables still exports", no_tables["tables"], 0)
check("and falls back to page text",
      load_workbook(os.path.join(OUT, "c_hard.xlsx")).sheetnames, ["Text"])

print("\n== powerpoint ==")
for pdf_font, family in [("BCDFEE+Cambria", "Cambria"),
                         ("ABCDEF+TimesNewRomanPSMT", "Times New Roman"),
                         ("Calibri-Italic", "Calibri")]:
    check("maps %s" % pdf_font, convert._pptx_font(pdf_font), family)

res = convert.to_pptx(fitz.open(SRC), os.path.join(OUT, "c_edit.pptx"))
deck = Presentation(os.path.join(OUT, "c_edit.pptx"))
slide = deck.slides[0]
pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
check("a background carries the page graphics", len(pictures), 1)
check("text is in real text boxes",
      len([sh for sh in slide.shapes if sh.has_text_frame]), lambda n: n >= 3)

runs = [r for sh in slide.shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs for r in p.runs]
check("every run names its font", all(r.font.name for r in runs), True)
bold = [r.text.strip() for r in runs if r.font.bold]
italic = [r.text.strip() for r in runs if r.font.italic]
check("bold survives inside a mixed line", any("proposed terms" == t for t in bold), True)
check("italic survives too", any("no later than" == t for t in italic), True)
check("colour is carried over",
      any(r.font.color and r.font.color.rgb is not None for r in runs), True)

mixed = convert.to_pptx(fitz.open(HARD), os.path.join(OUT, "c_hard.pptx"))
check("landscape pages are fitted, not squashed", mixed["resized_pages"],
      lambda n: n and n >= 2)
deck2 = Presentation(os.path.join(OUT, "c_hard.pptx"))
wide = [sh for sh in deck2.slides[1].shapes if sh.shape_type == 13][0]
check("a fitted page keeps its shape",
      round((wide.width / wide.height), 2), lambda r: 1.25 < r < 1.35)
check("and is centred on the slide", wide.top, lambda t: t > 0)

flat = convert.to_pptx(fitz.open(SRC), os.path.join(OUT, "c_img.pptx"), mode="image")
img_deck = Presentation(os.path.join(OUT, "c_img.pptx"))
check("picture mode has no text boxes",
      len([sh for sh in img_deck.slides[0].shapes if sh.has_text_frame]), 0)

txt = convert.to_pptx(fitz.open(SRC), os.path.join(OUT, "c_txt.pptx"), mode="text")
txt_deck = Presentation(os.path.join(OUT, "c_txt.pptx"))
check("text mode has no background picture",
      len([sh for sh in txt_deck.slides[0].shapes if sh.shape_type == 13]), 0)

print("\n" + "=" * 64)
print("No issues found." if not findings else "%d issue(s): %s" % (len(findings), findings))
sys.exit(1 if findings else 0)
