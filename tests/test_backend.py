"""End-to-end exercise of every backend feature against a real Word-exported PDF."""
import io
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pymupdf as fitz
from PIL import Image, ImageDraw

from backend import annots, convert, decorate, forms, pages, security, signatures, textedit
from backend.session import Session

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
SRC = os.path.join(OUT, "fixture.pdf")
DOCX = os.path.join(OUT, "fixture.docx")
results = []


def check(label, fn):
    try:
        detail = fn()
        results.append((True, label, detail))
        print("  PASS  %-34s %s" % (label, detail if detail else ""))
    except Exception as exc:
        results.append((False, label, repr(exc)))
        print("  FAIL  %-34s %r" % (label, exc))


def fresh():
    return fitz.open(SRC)


def make_signature_png(path):
    image = Image.new("RGB", (420, 140), "white")
    draw = ImageDraw.Draw(image)
    draw.line([(30, 100), (90, 40), (140, 105), (200, 45), (250, 100), (330, 60)],
              fill=(15, 20, 90), width=7)
    draw.line([(40, 118), (340, 118)], fill=(15, 20, 90), width=2)
    image.save(path)
    return path


print("\n== session ==")
session = Session()
check("open", lambda: session.open(SRC)["page_count"])
check("render", lambda: session.render(0, 1.0)["width"])
check("thumbnails", lambda: len(session.thumbnails(0, 1, 120)))
check("search", lambda: len(session.search("purchase")["hits"]))
check("outline round trip", lambda: (
    session.set_outline([{"level": 1, "title": "Terms", "page": 0}]),
    len(session.outline()))[1])
check("metadata", lambda: session.set_metadata({"title": "Memo", "author": "WCP"})["title"])


def undo_cycle():
    session.checkpoint()
    pages.insert_blank(session.doc, 1)
    grown = session.doc.page_count
    session.undo()
    return "%d -> %d" % (grown, session.doc.page_count)


check("undo/redo", undo_cycle)
check("save", lambda: os.path.basename(session.save(os.path.join(OUT, "session.pdf"))["name"]))

print("\n== pages ==")
check("rotate", lambda: pages.rotate(fresh(), [0], 90)["rotated"])
check("insert blank", lambda: pages.insert_blank(fresh(), 1)["page_count"])
check("duplicate", lambda: pages.duplicate(fresh(), [0])["page_count"])


def delete_page():
    d = fresh()
    pages.insert_blank(d, 1)
    return pages.delete(d, [1])["page_count"]


check("delete", delete_page)


def move_page():
    d = fresh()
    pages.insert_blank(d, 1)
    return pages.move(d, 1, 0)["page_count"]


check("move", move_page)
check("crop", lambda: pages.crop(fresh(), 0, [50, 50, 500, 700])["rect"][2])
check("extract", lambda: pages.extract(fresh(), [0], os.path.join(OUT, "extract.pdf"))["pages"])
check("merge", lambda: pages.merge(fresh(), [SRC])["page_count"])


def split_doc():
    d = fresh()
    pages.insert_blank(d, 1)
    return pages.split(d, os.path.join(OUT, "split"), "every", 1)["count"]


check("split", split_doc)

print("\n== annotations ==")


def markup_all():
    d = fresh()
    rect = d[0].search_for("purchase price")[0]
    out = []
    for kind in ("highlight", "underline", "strikeout", "squiggly"):
        out.append(annots.add_markup(d, 0, kind, [list(rect)])["type"])
    return ",".join(out)


check("markup (word snapped)", markup_all)
check("sticky note", lambda: annots.add_note(fresh(), 0, [300, 300], "Check this")["type"])
check("free text box", lambda: annots.add_textbox(fresh(), 0, [80, 600, 400, 640], "Draft")["type"])
check("shapes", lambda: ",".join(
    annots.add_shape(fresh(), 0, k, [[100, 500], [300, 560]])["type"]
    for k in ("rect", "circle", "line", "arrow")))
check("ink", lambda: annots.add_ink(fresh(), 0, [[[100, 400], [150, 430], [200, 400]]])["type"])
check("stamp", lambda: annots.add_stamp(fresh(), 0, [380, 80, 540, 140], "approved")["type"])


def annot_lifecycle():
    d = fresh()
    made = annots.add_note(d, 0, [200, 200], "temp")
    listed = annots.listing(d, 0)
    annots.update(d, 0, made["id"], content="updated")
    annots.delete(d, 0, made["id"])
    return "listed %d, deleted ok" % len(listed)


check("list/update/delete", annot_lifecycle)
check("flatten annots", lambda: (
    lambda d: (annots.add_note(d, 0, [200, 200], "x"), annots.flatten(d),
               len(annots.listing(d, 0)))[2])(fresh()))


def link_ops():
    d = fresh()
    annots.add_link(d, 0, [80, 700, 300, 720], uri="https://example.com")
    found = annots.links(d, 0)
    annots.delete_link(d, 0, 0)
    return "added %d, now %d" % (len(found), len(annots.links(d, 0)))


check("links", link_ops)

print("\n== images & signatures ==")
sig_png = make_signature_png(os.path.join(OUT, "sig.png"))
check("insert image", lambda: annots.add_image(fresh(), 0, [80, 640, 260, 700], sig_png)["ok"])

created = []


def sig_add():
    entry = signatures.add(sig_png, "Test Signer", "Managing Partner")
    created.append(entry["id"])
    return entry["id"]


check("signature add", sig_add)
check("signature list", lambda: len(signatures.listing()))
check("signature rename", lambda: signatures.rename(created[0], "Test Signer", "Partner")["role"])


def sig_place_flat():
    d = fresh()
    signatures.place(d, 0, created[0], [140, 655, 330, 700])
    # A placed signature must be page content, not a removable annotation.
    return "annots on page: %d" % len(annots.listing(d, 0))


check("signature placed flat", sig_place_flat)
check("signature transparent bg", lambda: (
    lambda p: Image.open(p).convert("RGBA").getextrema()[3][0])(
        os.path.join(signatures.SIG_DIR, signatures._load()[-1]["file"])))

print("\n== forms ==")


def form_cycle():
    d = fresh()
    forms.add_field(d, 0, "text", [350, 470, 540, 492], "signer_name")
    forms.add_field(d, 0, "checkbox", [350, 500, 366, 516], "agreed")
    forms.add_field(d, 0, "dropdown", [350, 524, 540, 546], "status",
                    options=["Draft", "Final"])
    data = d.tobytes()
    d2 = fitz.open(stream=data, filetype="pdf")
    forms.set_value(d2, 0, "signer_name", "Ernie Willmore")
    forms.set_value(d2, 0, "status", "Final")
    listed = forms.listing(d2)
    csv_path = os.path.join(OUT, "fields.csv")
    forms.export_csv(d2, csv_path)
    forms.import_csv(d2, csv_path)
    forms.flatten(d2)
    return "%d fields, csv %d bytes" % (len(listed), os.path.getsize(csv_path))


check("create/fill/export/flatten", form_cycle)

print("\n== security ==")
check("describe", lambda: security.describe(fresh())["encrypted"])


def protect_cycle():
    d = fresh()
    path = os.path.join(OUT, "protected.pdf")
    security.save_protected(d, path, user_password="open123", allowed=["print"])
    locked = fitz.open(path)
    needs = locked.needs_pass
    locked.authenticate("open123")
    plain = os.path.join(OUT, "unlocked.pdf")
    security.save_unprotected(locked, plain)
    return "needs_pass=%s, unlocked=%s" % (needs, not fitz.open(plain).needs_pass)


check("protect / unprotect", protect_cycle)


def redact_cycle():
    d = fresh()
    security.mark_matches(d, "4,250,000")
    security.apply(d)
    text = fitz.open(stream=d.tobytes(), filetype="pdf")[0].get_text()
    return "removed=%s" % ("4,250,000" not in text)


check("redaction removes text", redact_cycle)

print("\n== conversion ==")
check("to text", lambda: os.path.getsize(convert.to_text(fresh(), os.path.join(OUT, "out.txt"))["path"]))
check("to images", lambda: convert.to_images(fresh(), os.path.join(OUT, "images"), 120)["count"])
check("to xlsx", lambda: convert.to_xlsx(fresh(), os.path.join(OUT, "out.xlsx"))["tables"])
check("to pptx (editable)", lambda: convert.to_pptx(fresh(), os.path.join(OUT, "out.pptx"))["slides"])
check("to pptx (image)", lambda: convert.to_pptx(fresh(), os.path.join(OUT, "out_img.pptx"), "image")["slides"])
check("to docx", lambda: os.path.getsize(convert.to_docx(fresh(), os.path.join(OUT, "out.docx"))["path"]))


def docx_roundtrip():
    """The edit -> Word path is the one that mattered most, so check the text."""
    d = fresh()
    span = [s for b in textedit.layout(d[0])["blocks"] for ln in b["lines"]
            for s in ln["spans"] if "4,250,000" in s["text"]][0]
    textedit.edit_span(d, span["id"], span["text"].replace("$4,250,000", "$9,875,400"))
    reloaded = fitz.open(stream=d.tobytes(), filetype="pdf")
    path = os.path.join(OUT, "edited.docx")
    convert.to_docx(reloaded, path)
    from docx import Document
    body = "\n".join(p.text for p in Document(path).paragraphs)
    tables = " ".join(c.text for t in Document(path).tables for r in t.rows for c in r.cells)
    assert "9,875,400" in body, "edited value missing from Word output"
    assert " " not in body, "non-breaking spaces leaked into Word output"
    assert "Deposit" in tables, "table lost in Word output"
    return "edit survived to Word, table intact"


check("edit -> Word round trip", docx_roundtrip)
check("office -> pdf", lambda: convert.from_files([DOCX]).page_count)
check("images -> pdf", lambda: convert.from_files([sig_png]).page_count)
check("compress", lambda: convert.compress(fresh(), os.path.join(OUT, "small.pdf"), "high")["size"])
check("ocr engine is built in", lambda: convert.ocr_status()["available"],)

print("\n== decoration ==")
check("watermark text", lambda: decorate.watermark_text(fresh(), "CONFIDENTIAL")["pages"])
check("watermark image", lambda: decorate.watermark_image(fresh(), sig_png)["pages"])
check("page numbers", lambda: decorate.stamp_text(fresh(), "Page {page} of {pages}")["pages"])
check("header", lambda: decorate.stamp_text(fresh(), "WCP", "top-right")["pages"])
check("background", lambda: decorate.background(fresh(), color=(1, 1, 0.9))["pages"])


def stamp_text_is_clean():
    d = fresh()
    decorate.stamp_text(d, "Page {page} of {pages}")
    text = fitz.open(stream=d.tobytes(), filetype="pdf")[0].get_text()
    assert "Page 1 of 1" in text, "stamped text did not extract correctly"
    return "extracts as typed"


check("stamped text extracts", stamp_text_is_clean)

for identifier in created:
    try:
        signatures.remove(identifier)
    except Exception:
        pass

bad = [r for r in results if not r[0]]
print("\n%d checks, %d failed" % (len(results), len(bad)))
for _, label, detail in bad:
    print("  FAILED:", label, detail)
print("RESULT:", "PASS" if not bad else "FAIL")
