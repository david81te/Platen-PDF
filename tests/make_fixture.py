"""Build a Word-exported PDF fixture (mirrors the real-world input: .docx -> PDF)."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches

OUT = os.path.join(os.path.dirname(__file__), "output")
os.makedirs(OUT, exist_ok=True)
docx_path = os.path.join(OUT, "fixture.docx")
pdf_path = os.path.join(OUT, "fixture.pdf")

d = Document()
d.add_heading("Willmore Capital Partners", level=1)
p = d.add_paragraph("This memorandum summarizes the ")
p.add_run("proposed terms").bold = True
p.add_run(" of the transaction dated January 5, 2026. The parties intend to close ")
p.add_run("no later than").italic = True
p.add_run(" the end of the quarter, subject to customary conditions.")

d.add_paragraph(
    "The purchase price is $4,250,000, payable in cash at closing. Seller shall deliver "
    "clear title free of encumbrances. Buyer has completed its diligence review and waives "
    "the financing contingency described in Section 4.2 of the letter of intent."
)
t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
for i, row in enumerate([["Item", "Amount", "Due"],
                         ["Deposit", "$100,000", "Jan 15"],
                         ["Balance", "$4,150,000", "Mar 31"]]):
    for j, val in enumerate(row):
        t.rows[i].cells[j].text = val
d.add_paragraph()
d.add_paragraph("Signature: ______________________")
d.save(docx_path)

import win32com.client as win32
word = win32.Dispatch("Word.Application")
word.Visible = False
try:
    doc = word.Documents.Open(docx_path)
    doc.ExportAsFixedFormat(pdf_path, 17)  # 17 = wdExportFormatPDF
    doc.Close(False)
finally:
    word.Quit()
print("wrote", pdf_path, os.path.getsize(pdf_path), "bytes")
